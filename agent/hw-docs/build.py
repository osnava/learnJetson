#!/usr/bin/env python3
"""docling corpus builder (issue #28) — docling conversion wired end to end.

One invocation = one fetch.sh pass over a set of sources:

    convert (docling, standard pipeline)  ->  JSON source of truth (core
    docs only) + raw markdown; chunk (HybridChunker, markdown-serialized,
    heading breadcrumbs + page provenance from the DoclingDocument);
    embed (small local model, GPU when available) -> per-doc index shards
    beside the corpus; derive the corpus-confirmed wrap table.

Substrate decisions (issue #28, from the #27 spike):

- JSON is the source of truth, md is the rendering agents read. The TRM
  skips the JSON (a full one would be ~970 MB); it converts in page-range
  slabs so memory stays bounded and the grind is resumable — each slab's
  index shard is written atomically and skipped on rerun.
- Page provenance comes from object-level `prov.page_no` (verified
  absolute under `page_range`), never from text anchors.
- Every text transform lives in normalize.py (shared with search.py and
  the #29 grader); this file only orchestrates.
- The xlsx->csv pinmux path is carried over unchanged from the v1 pymupdf
  era (explicitly out of the rebuild's hard wall; v1 files deleted in
  #30's sweep).

Run under a docling-capable Python (fetch.sh resolves one; see README.md).
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import embed  # noqa: E402
import normalize  # noqa: E402

MODEL = embed.MODEL                    # chunker tokenizer == embedder == window
MAX_CHUNK_TOKENS = embed.MAX_TOKENS

# Docling imports are lazy: tests (and search-side tooling) import this
# module on machines without docling; only the conversion paths need it.
_DOCLING = None


def _docling():
    """Import and cache the docling API; die with guidance if absent."""
    global _DOCLING
    if _DOCLING is None:
        try:
            from docling.document_converter import DocumentConverter, PdfFormatOption
            from docling.datamodel.base_models import InputFormat
            from docling.datamodel.document import ConversionStatus
            from docling_core.transforms.chunker import HybridChunker
            from docling_core.transforms.serializer.base import BaseSerializerProvider
            from docling_core.transforms.serializer.markdown import MarkdownDocSerializer
            from docling_core.types.doc import ImageRefMode
        except ImportError as e:
            sys.exit(f"docling not importable under {sys.executable}: {e}\n"
                     "fetch.sh resolves a docling-capable Python automatically; "
                     "by hand:  uv tool install docling   (or pip install docling)")
        _DOCLING = dict(DocumentConverter=DocumentConverter,
                        PdfFormatOption=PdfFormatOption, InputFormat=InputFormat,
                        ConversionStatus=ConversionStatus,
                        HybridChunker=HybridChunker,
                        BaseSerializerProvider=BaseSerializerProvider,
                        MarkdownDocSerializer=MarkdownDocSerializer,
                        ImageRefMode=ImageRefMode)
    return _DOCLING


# ---------------------------------------------------------------- xlsx path
# Carried over verbatim from the v1 pymupdf era — issue #28 kept the
# pinmux xlsx->csv path unchanged; it is deliberately not part of the
# rebuild.

def _safe(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("_") or "sheet"


def convert_xlsx(src: Path, dst: Path) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(str(src), read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets:
        csv_path = dst.parent / f"{src.stem}.{_safe(ws.title)}.csv"
        with csv_path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                if any(c is not None for c in row):
                    writer.writerow(["" if c is None else c for c in row])
        rows.append((ws.title, csv_path.name))
    listing = "\n".join(f"- `{csv}` - sheet '{title}'" for title, csv in rows)
    dst.write_text(
        f"# {src.stem} (XLSX)\n\n"
        "Pinmux workbook - each sheet exported as CSV (one row per pin, grep-friendly):\n\n"
        f"{listing}\n",
        encoding="utf-8",
    )


# ------------------------------------------------------------- pdf pipeline

def pdf_page_count(src: Path) -> int:
    import pypdfium2 as pdfium

    return len(pdfium.PdfDocument(str(src)))


def convert_slab(src: Path, first: int, last: int):
    """Standard docling pipeline over one page range; images stay as
    placeholders so md never grows image payloads (spike decision).
    page_range goes to convert() (docling 2.4+ moved it off the pipeline
    options) and prov page_no stays absolute under it (spike-verified:
    pp.120-139 came back as 120-139)."""
    d = _docling()
    conv = d["DocumentConverter"](format_options={
        d["InputFormat"].PDF: d["PdfFormatOption"]()})
    res = conv.convert(str(src), page_range=(first, last))
    if res.status != d["ConversionStatus"].SUCCESS:
        raise RuntimeError(f"docling failed on {src.name} pp.{first}-{last}: {res.status}")
    return res.document


_MD_CHUNKS = None


def _md_chunks_cls():
    """Build (once) the chunk-serializer provider class: chunks serialize
    as markdown — table rows stay rows instead of the default key-value
    linearization, so retrieved snippets stay greppable."""
    global _MD_CHUNKS
    if _MD_CHUNKS is None:
        d = _docling()

        class _MdChunks(d["BaseSerializerProvider"]):
            def get_serializer(self, doc):
                return d["MarkdownDocSerializer"](doc=doc)

        _MD_CHUNKS = _MdChunks
    return _MD_CHUNKS


def chunk_records(doc, stem: str) -> list[dict]:
    """Provenance-carrying chunk records — the retrieval contract:

        {"doc": stem, "headings": [...], "pages": [28], "text": "..."}
    """
    chunker = _docling()["HybridChunker"](
        tokenizer=MODEL, max_tokens=MAX_CHUNK_TOKENS,
        serializer_provider=_md_chunks_cls()(), repeat_table_header=True)
    records = []
    for chunk in chunker.chunk(doc):
        pages = sorted({p.page_no for item in chunk.meta.doc_items
                        for p in (item.prov or [])})
        records.append({
            "doc": stem,
            "headings": [normalize.canonical(h) for h in (chunk.meta.headings or [])],
            "pages": pages,
            "text": normalize.canonical(chunk.text),
        })
    return records


def _atomic_write(path: Path, data: str) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(data, encoding="utf-8", newline="\n")
    tmp.replace(path)


def write_shard(index_dir: Path, shard: str, records: list[dict], emb, meta: dict) -> None:
    import numpy

    _atomic_write(index_dir / f"{shard}.chunks.jsonl",
                  "".join(json.dumps(r, ensure_ascii=False) + "\n" for r in records))
    tmp = index_dir / f"{shard}.emb.npy.tmp"
    with open(tmp, "wb") as f:  # numpy.save would append .npy to the tmp name
        numpy.save(f, emb)
    tmp.replace(index_dir / f"{shard}.emb.npy")
    _atomic_write(index_dir / f"{shard}.meta.json",
                  json.dumps(meta, ensure_ascii=False, indent=1))


def build_pdf(src: Path, stem: str, md_dir: Path, index_dir: Path,
              slab_pages: int | None, persist_json: bool,
              log=print) -> tuple[list[Path], bool]:
    """Convert one PDF. Returns (raw-md part files, all slabs complete) —
    the md rendering is only assembled from a COMPLETE part set, so an
    interrupted grind can never leave a partial md behind that staleness
    checks would mistake for a finished document."""
    n_pages = pdf_page_count(src)
    slabs = ([(a, min(a + slab_pages - 1, n_pages))
              for a in range(1, n_pages + 1, slab_pages)]
             if slab_pages else [(1, n_pages)])
    slab_mode = len(slabs) > 1
    parts: list[Path] = []

    for first, last in slabs:
        shard = f"{stem}.p{first}-{last}" if slab_mode else stem
        chunks_f = index_dir / f"{shard}.chunks.jsonl"
        emb_f = index_dir / f"{shard}.emb.npy"
        part_f = md_dir / f"{shard}.raw.md"
        if chunks_f.is_file() and emb_f.is_file() and part_f.is_file():
            log(f"   {stem} pp.{first}-{last}: shard cached")
            parts.append(part_f)
            continue

        t0 = time.time()
        doc = convert_slab(src, first, last)
        raw_md = doc.export_to_markdown(
            image_mode=_docling()["ImageRefMode"].PLACEHOLDER)
        # part before shard: a crash between the two re-converts cleanly,
        # while the reverse order would strand a shard with no raw text
        _atomic_write(part_f, raw_md)
        records = chunk_records(doc, stem)
        if not records:
            log(f"   !! {stem} pp.{first}-{last}: no chunks — skipped")
            continue
        texts = [r["text"] for r in records]
        emb = embed.embed_texts(texts)
        write_shard(index_dir, shard, records, emb, {
            "doc": stem, "model": MODEL, "dim": int(emb.shape[1]),
            "chunks": len(records), "pages": f"{first}-{last}",
            "pdf": src.name, "built": time.strftime("%Y-%m-%dT%H:%M:%S"),
        })
        if persist_json and not slab_mode:
            doc.save_as_json(md_dir / f"{stem}.json")
        parts.append(part_f)
        log(f"   {stem} pp.{first}-{last}/{n_pages}: {len(records)} chunks, "
            f"{time.time() - t0:.0f}s")
    return parts, len(parts) == len(slabs)


def finalize_md(md_dir: Path, index_dir: Path, stem: str, parts: list[Path],
                wrap_table: list[dict], log=print) -> None:
    """Concatenate raw parts -> normalized md rendering; parts removed."""
    raw = "".join(p.read_text(encoding="utf-8", errors="replace") for p in parts)
    _atomic_write(md_dir / f"{stem}.md", normalize.render(raw, wrap_table))
    for p in parts:
        p.unlink(missing_ok=True)
    log(f"   {stem}: md rendering written ({len(raw) // 1024} KB raw)")


def merge_wrap_table(index_dir: Path, new_texts: list[str]) -> list[dict]:
    existing = normalize.load_wrap_table(index_dir / "wrap_table.json")
    merged = {(e["wrapped"], e["joined"]) for e in existing}
    merged |= {(e["wrapped"], e["joined"]) for e in normalize.derive_wrap_table(new_texts)}
    return [{"wrapped": w, "joined": j}
            for w, j in sorted(merged, key=lambda wj: -len(wj[0]))]


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--md-dir", type=Path, default=Path("md"))
    ap.add_argument("--index-dir", type=Path, default=Path("md/index"))
    ap.add_argument("--persist-json", action="store_true",
                    help="keep the JSON source of truth (core docs; the TRM "
                         "would be ~970 MB)")
    ap.add_argument("--slab-pages", type=int, default=None,
                    help="convert in N-page slabs (resumable; use for the TRM)")
    ap.add_argument("sources", nargs="+", metavar="SRC",
                    help="pdf or xlsx; rename with SRC=STEM (schematics)")
    args = ap.parse_args(argv)

    args.md_dir.mkdir(parents=True, exist_ok=True)
    args.index_dir.mkdir(parents=True, exist_ok=True)

    parsed = []
    for s in args.sources:
        path, _, stem = s.partition("=")
        src = Path(path)
        if not src.is_file():
            print(f"!! missing source: {src}", file=sys.stderr)
            return 1
        parsed.append((src, stem or src.stem))

    failures = 0
    parts_by_stem: dict[str, tuple[list[Path], bool]] = {}
    raw_texts: list[str] = []
    for src, stem in parsed:
        print(f">> {src.name}" + (f" (as {stem})" if stem != src.stem else ""))
        try:
            if src.suffix.lower() == ".xlsx":
                convert_xlsx(src, args.md_dir / f"{stem}.md")
                print(f"   {stem}: xlsx -> md + per-sheet CSVs")
            else:
                parts, complete = build_pdf(src, stem, args.md_dir,
                                            args.index_dir, args.slab_pages,
                                            args.persist_json)
                prev = parts_by_stem.get(stem, ([], True))
                parts_by_stem[stem] = (prev[0] + parts, prev[1] and complete)
                raw_texts.extend(p.read_text(encoding="utf-8", errors="replace")
                                 for p in parts)
        except Exception as e:  # one bad source must not sink the batch
            print(f"   !! {stem} failed: {e}", file=sys.stderr)
            failures += 1

    wrap_table = merge_wrap_table(args.index_dir, raw_texts)
    _atomic_write(args.index_dir / "wrap_table.json",
                  json.dumps(wrap_table, ensure_ascii=False, indent=1))
    for stem, (parts, complete) in parts_by_stem.items():
        if not parts:
            continue
        if complete:
            finalize_md(args.md_dir, args.index_dir, stem, parts, wrap_table)
        else:
            print(f"   !! {stem}: md NOT written - slabs incomplete; "
                  "re-run to resume (finished slabs are kept)", file=sys.stderr)
            failures += 1
    print(f"wrap table: {len(wrap_table)} entries -> "
          f"{args.index_dir / 'wrap_table.json'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
